from datetime import datetime

import pytest
from openpyxl import Workbook

from dva.agents.loaders import (
    detect_header_row,
    load_csv_bytes,
    load_files,
    normalize_name,
    parse_html_tables,
)

VOCAB = {normalize_name(w) for w in ["store_id", "sku", "quantity", "transaction_date",
                                     "unit_price", "region", "Store_No"]}  # fmt: skip


def test_csv_fixture_keeps_every_cell_as_text(fixtures_dir):
    result = load_files([fixtures_dir / "type_errors" / "sales.csv"], VOCAB)
    (table,) = result.tables
    assert result.issues == []
    assert len(table.frame) == 2000 and table.header_row == 1
    assert all(table.frame[c].map(type).eq(str).all() for c in table.frame.columns)
    assert table.frame["unit_price"].str.startswith("$").any()


def test_csv_delimiter_and_encoding_are_detected():
    data = "Store_No;Region\nS1;Nörth\nS2;South\n".encode("cp1252")
    table = load_csv_bytes(data, "stores.csv")
    assert list(table.frame.columns) == ["Store_No", "Region"]
    assert table.frame["Region"].tolist() == ["Nörth", "South"]


def test_excel_fixture_selects_data_sheet_and_header_row(fixtures_dir):
    result = load_files([fixtures_dir / "excel_messy" / "sales.xlsx"], VOCAB)
    (table,) = result.tables  # the Summary sheet has no header row and is skipped
    assert table.sheet == "Sales" and table.header_row == 3
    assert table.source == "sales.xlsx#Sales"
    assert len(table.frame) == 3000
    assert list(table.frame.columns)[:2] == ["transaction_date", "store_id"]
    assert isinstance(table.frame["transaction_date"].iloc[0], datetime)


def test_excel_with_no_recognizable_sheet_is_an_issue(tmp_path):
    workbook = Workbook()
    workbook.active.append(["just", "notes"])
    path = tmp_path / "notes.xlsx"
    workbook.save(path)
    result = load_files([path], VOCAB)
    assert result.tables == [] and "no sheet" in result.issues[0].detail


def test_email_fixture_attachment_and_html_table(fixtures_dir):
    result = load_files([fixtures_dir / "email_table" / f for f in ("sales.eml", "stores.eml")],
                        VOCAB)  # fmt: skip
    sales, stores = result.tables
    assert sales.source == "sales.eml#sales_export.csv" and len(sales.frame) == 500
    assert stores.source == "stores.eml#html-table-1" and len(stores.frame) == 40
    assert list(stores.frame.columns) == ["store_id", "region"]


def test_unsupported_and_unreadable_files_become_issues(tmp_path):
    (tmp_path / "data.parquet").write_bytes(b"PAR1")
    (tmp_path / "broken.xlsx").write_bytes(b"not a zip")
    result = load_files([tmp_path / "data.parquet", tmp_path / "broken.xlsx"], VOCAB)
    assert result.tables == []
    assert [i.source for i in result.issues] == ["data.parquet", "broken.xlsx"]
    assert "unsupported" in result.issues[0].detail


@pytest.mark.parametrize(
    ("rows", "expected"),
    [
        ([["Report"], [], ["Store", "SKU", "Qty"], ["S1", "A", 1]], None),
        ([["Report"], ["store_id", "sku", "quantity"], ["S1", "A", 1]], 1),
        ([["store id", "SKU"], ["S1", "A"]], 0),
    ],  # fmt: skip
)
def test_header_row_detection(rows, expected):
    assert detect_header_row(rows, VOCAB) == expected


def test_html_table_parser_handles_entities_and_whitespace():
    html = (
        "<table><tr><th>store_id</th><th>region</th></tr>"
        "<tr><td> S1 </td><td>A&amp;B</td></tr></table>"
    )
    assert parse_html_tables(html) == [[["store_id", "region"], ["S1", "A&B"]]]
