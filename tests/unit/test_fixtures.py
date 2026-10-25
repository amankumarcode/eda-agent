"""Fixture generation is deterministic, matches the committed files, and the injected
issues in the data agree with ground_truth.json (checked independently of the generator)."""

import io
import json
import re
from email import message_from_bytes, policy
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from dva.state import IssueType

FIXTURES = [
    "clean", "renamed_columns", "missing_field", "unit_mismatch", "implausible_values",
    "category_variants", "join_failure", "type_errors", "duplicates", "excel_messy",
    "email_table", "kitchen_sink",
]  # fmt: skip
EXPECTED_SALES_ROWS = {
    "clean": 2000, "renamed_columns": 2000, "missing_field": 2000, "unit_mismatch": 5000,
    "implausible_values": 5000, "category_variants": 3000, "join_failure": 4000,
    "type_errors": 2000, "duplicates": 2000, "excel_messy": 3000, "email_table": 500,
    "kitchen_sink": 10000,
}  # fmt: skip


def _tree_bytes(root: Path) -> dict[str, bytes]:
    return {
        str(p.relative_to(root)): p.read_bytes() for p in sorted(root.rglob("*")) if p.is_file()
    }


@pytest.fixture(scope="module")
def two_runs(generator, retail_model, tmp_path_factory):
    first, second = tmp_path_factory.mktemp("gen1"), tmp_path_factory.mktemp("gen2")
    generator.generate_all(first, retail_model)
    generator.generate_all(second, retail_model)
    return _tree_bytes(first), _tree_bytes(second)


def test_generation_is_deterministic(two_runs):
    first, second = two_runs
    assert first.keys() == second.keys()
    assert [k for k in first if first[k] != second[k]] == []


def test_committed_fixtures_match_generator(two_runs, fixtures_dir):
    generated, _ = two_runs
    committed = _tree_bytes(fixtures_dir)
    assert committed.keys() == generated.keys()
    assert [k for k in committed if committed[k] != generated[k]] == []


def test_every_fixture_present():
    assert set(FIXTURES) == set(__import__("generate_fixtures").BUILDERS)


def _gt(fixtures_dir, name):
    return json.loads((fixtures_dir / name / "ground_truth.json").read_text())


def _csv(fixtures_dir, name, table):
    return pd.read_csv(fixtures_dir / name / f"{table}.csv", dtype=str, keep_default_na=False)


@pytest.mark.parametrize("name", FIXTURES)
def test_ground_truth_is_consistent(fixtures_dir, retail_model, name):
    gt = _gt(fixtures_dir, name)
    assert gt["rows"]["sales"] == EXPECTED_SALES_ROWS[name]
    fields = retail_model.field_names()
    assert all(v is None or v in fields for v in gt["expected_field_mapping"].values())
    for item in gt["expected_issues"]:
        IssueType(item["issue_type"])
        assert ("count" in item) != ("count_approx" in item)
    columns: set[str] = set()
    for table, filename in gt["files"].items():
        assert (fixtures_dir / name / filename).is_file()
        if filename.endswith(".csv"):
            frame = _csv(fixtures_dir, name, table)
            assert len(frame) == gt["rows"][table]
            columns |= set(frame.columns)
    if all(f.endswith(".csv") for f in gt["files"].values()):
        assert columns == set(gt["expected_field_mapping"])


def _counts(gt):
    return {(i["field"], i["issue_type"]): i.get("count", i.get("count_approx"))
            for i in gt["expected_issues"]}  # fmt: skip


def test_clean_has_no_issues(fixtures_dir):
    sales = _csv(fixtures_dir, "clean", "sales")
    products, stores = (
        _csv(fixtures_dir, "clean", "products"),
        _csv(fixtures_dir, "clean", "stores"),
    )
    assert _gt(fixtures_dir, "clean")["expected_issues"] == []
    assert not sales.duplicated().any()
    assert sales["sku"].isin(products["sku"]).all()
    assert sales["store_id"].isin(stores["store_id"]).all()
    price, qty = sales["unit_price"].astype(float), sales["quantity"].astype(int)
    assert price.between(0.01, 500).all() and qty.between(-500, 5000).all()
    assert (qty * price).between(-10000, 50000).all()
    assert set(sales["channel"]) == {"store", "online", "marketplace"}
    assert pd.to_datetime(sales["transaction_date"], format="%Y-%m-%d").between(
        "2015-01-01", "2026-01-01").all()  # fmt: skip
    assert products["pack_size"].astype(int).between(1, 144).all()


def test_implausible_values_count(fixtures_dir):
    sales = _csv(fixtures_dir, "implausible_values", "sales")
    assert (sales["unit_price"].astype(float) > 500).sum() == 37


def test_join_failure_count(fixtures_dir):
    sales, products = (_csv(fixtures_dir, "join_failure", t) for t in ("sales", "products"))
    assert len(products) == 300
    assert (~sales["sku"].isin(products["sku"])).sum() == 412


def test_duplicates_count(fixtures_dir):
    assert _csv(fixtures_dir, "duplicates", "sales").duplicated().sum() == 50


def test_category_variants(fixtures_dir):
    channel = _csv(fixtures_dir, "category_variants", "sales")["channel"]
    assert (channel == "ZZ").sum() == 25
    assert set(channel) == {"Store", "B&M", "1", "online", "marketplace", "ZZ"}


def test_unit_mismatch_is_about_12x_in_case_stores(fixtures_dir):
    gt = _gt(fixtures_dir, "unit_mismatch")
    sales = _csv(fixtures_dir, "unit_mismatch", "sales")
    qty = sales["quantity"].astype(int)
    in_case = sales["store_id"].isin(gt["case_stores"])
    assert in_case.sum() == _counts(gt)[("quantity", "unit_mismatch")]
    assert 0.07 <= in_case.mean() <= 0.09
    ratio = qty[in_case].median() / qty[~in_case].median()
    assert 10 <= ratio <= 14


def test_type_errors_counts(fixtures_dir):
    counts = _counts(_gt(fixtures_dir, "type_errors"))
    sales = _csv(fixtures_dir, "type_errors", "sales")
    iso = sales["transaction_date"].str.fullmatch(r"\d{4}-\d{2}-\d{2}")
    plain = sales["unit_price"].str.fullmatch(r"\d+\.\d{2}")
    assert (~iso).sum() == counts[("transaction_date", "type_mismatch")]
    assert (~plain).sum() == counts[("unit_price", "type_mismatch")]
    assert sales["unit_price"].str.startswith("$").any()
    assert sales["unit_price"].str.contains(",").any()


def test_missing_field_has_no_channel(fixtures_dir):
    assert "channel" not in _csv(fixtures_dir, "missing_field", "sales").columns


def test_excel_layout(fixtures_dir):
    workbook = load_workbook(fixtures_dir / "excel_messy" / "sales.xlsx", read_only=True)
    assert workbook.sheetnames == ["Summary", "Sales"]
    rows = list(workbook["Sales"].iter_rows(values_only=True))
    assert rows[2][:2] == ("transaction_date", "store_id")
    assert len(rows) == 3 + 3000


def test_email_fixtures_parse(fixtures_dir):
    sales_msg = message_from_bytes((fixtures_dir / "email_table" / "sales.eml").read_bytes(),
                                   policy=policy.default)  # fmt: skip
    attachment = next(sales_msg.iter_attachments())
    assert attachment.get_filename() == "sales_export.csv"
    frame = pd.read_csv(io.StringIO(attachment.get_content()))
    assert len(frame) == 500
    stores_msg = message_from_bytes((fixtures_dir / "email_table" / "stores.eml").read_bytes(),
                                    policy=policy.default)  # fmt: skip
    html = stores_msg.get_body(preferencelist=("html",)).get_content()
    assert len(re.findall(r"<tr>", html)) == 1 + 40


def test_kitchen_sink_counts(fixtures_dir):
    gt = _gt(fixtures_dir, "kitchen_sink")
    counts = _counts(gt)
    sales, products = (_csv(fixtures_dir, "kitchen_sink", t) for t in ("sales", "products"))
    assert (
        "Cashier_Notes" in sales.columns and gt["expected_field_mapping"]["Cashier_Notes"] is None
    )
    assert (sales["Sell_Price"].astype(float) > 500).sum() == counts[
        ("unit_price", "implausible_value")
    ]
    assert (sales["Sales_Channel"] == "ZZ").sum() == counts[("channel", "unmapped_category")]
    assert (~sales["Item_No"].isin(products["UPC"])).sum() == counts[("sku", "join_failure")]
    assert sales.duplicated().sum() == counts[("*", "duplicate_rows")]
    in_case = sales["Store_No"].isin(gt["case_stores"]).sum()
    assert in_case == counts[("quantity", "unit_mismatch")]
