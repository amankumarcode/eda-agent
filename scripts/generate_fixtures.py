"""Generate test fixtures with known, injected issues, plus ground_truth.json per fixture.

Usage: python scripts/generate_fixtures.py [--out DIR]

Seeded (seed 42) and byte-for-byte deterministic. Rerun only when the data model changes;
the fixtures are committed so tests are reproducible.

Ground-truth conventions:
- expected_field_mapping is flat (source column -> model field or null) across all tables.
- count is the number of affected rows; for missing_field and unmapped_column that is
  every row of the table. duplicate_rows uses field "*" (the whole row).
- count_approx allows +/-10% where detection is statistical.
"""

from __future__ import annotations

import argparse
import io
import json
import re
import shutil
import zipfile
import zlib
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook

from dva.config import load_settings
from dva.data_model import DataModel, make_provider

SEED = 42
N_STORES = 40
N_PRODUCTS = 300
REGIONS = ["North", "South", "East", "West"]
CATEGORIES = ["Beverages", "Snacks", "Dairy", "Household", "Personal Care", "Frozen"]
PACK_SIZES = [1, 6, 12, 12, 12, 24]
CHANNELS = ["store", "online", "marketplace"]
CHANNEL_P = [0.7, 0.2, 0.1]
DATE_START = date(2023, 1, 1)
DATE_DAYS = 730
CASE_PACK = 12
FIXED_TIMESTAMP = datetime(2024, 1, 1, 0, 0, 0)
ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)
EMAIL_BOUNDARY = "dva-fixture-boundary"
CASHIER_NOTES = ["", "", "", "price override", "loyalty card", "manager ok", "void retry"]


@dataclass
class Ctx:
    """Everything a fixture builder needs: its own RNG, output folder, and master data."""

    name: str
    rng: np.random.Generator
    out_dir: Path
    model: DataModel
    stores: pd.DataFrame
    products: pd.DataFrame
    prices: np.ndarray


# --- master data and sales -------------------------------------------------------------


def make_stores(rng: np.random.Generator) -> pd.DataFrame:
    """Return the 40-store master with a region per store."""
    ids = [f"S{i:03d}" for i in range(1, N_STORES + 1)]
    return pd.DataFrame({"store_id": ids, "region": rng.choice(REGIONS, N_STORES)})


def make_products(rng: np.random.Generator) -> tuple[pd.DataFrame, np.ndarray]:
    """Return the product master and a base unit price per product (same order)."""
    products = pd.DataFrame(
        {
            "sku": [f"SKU{i:05d}" for i in range(1, N_PRODUCTS + 1)],
            "category": rng.choice(CATEGORIES, N_PRODUCTS),
            "pack_size": rng.choice(PACK_SIZES, N_PRODUCTS),
        }
    )
    prices = rng.lognormal(np.log(4.0), 0.7, N_PRODUCTS).clip(0.5, 60.0).round(2)
    return products, prices


def make_sales(ctx: Ctx, n: int, store_p: np.ndarray | None = None) -> pd.DataFrame:
    """Clean sales rows: valid dates, known stores and SKUs, plausible values, no duplicates.

    Oversamples by 2%, drops chance duplicates, then keeps a random n rows.
    """
    rng, n_draw = ctx.rng, int(n * 1.02) + 10
    sales = _draw_sales(ctx, n_draw, store_p).drop_duplicates()
    sales = sales.iloc[np.sort(rng.choice(len(sales), n, replace=False))]
    sales = sales.sort_values(["transaction_date", "store_id", "sku"], kind="stable")
    return sales.reset_index(drop=True)


def _draw_sales(ctx: Ctx, n: int, store_p: np.ndarray | None) -> pd.DataFrame:
    rng = ctx.rng
    sku_idx = rng.integers(0, len(ctx.products), n)
    qty = rng.poisson(3, n) + 1
    qty = np.where(rng.random(n) < 0.02, -rng.integers(1, 4, n), qty)
    days = rng.integers(0, DATE_DAYS, n)
    sales = pd.DataFrame(
        {
            "transaction_date": [(DATE_START + timedelta(days=int(d))).isoformat() for d in days],
            "store_id": rng.choice(ctx.stores["store_id"].to_numpy(), n, p=store_p),
            "sku": ctx.products["sku"].to_numpy()[sku_idx],
            "quantity": qty,
            "unit_price": ctx.prices[sku_idx],
            "channel": rng.choice(CHANNELS, n, p=CHANNEL_P),
            "promo_flag": rng.random(n) < 0.15,
        }
    )
    return sales


def case_store_weights(ctx: Ctx, case_stores: list[str], share: float) -> np.ndarray:
    """Store probabilities where the case stores together get ``share`` of rows."""
    is_case = ctx.stores["store_id"].isin(case_stores).to_numpy()
    weights = np.where(is_case, share / is_case.sum(), (1 - share) / (~is_case).sum())
    return weights / weights.sum()


# --- injections -------------------------------------------------------------------------


def inject_cases(sales: pd.DataFrame, case_stores: list[str]) -> int:
    """Multiply quantity by the case pack in the case stores. Returns affected rows."""
    mask = sales["store_id"].isin(case_stores)
    sales.loc[mask, "quantity"] = sales.loc[mask, "quantity"] * CASE_PACK
    return int(mask.sum())


def price_eligible(sales: pd.DataFrame) -> pd.Series:
    """Rows whose price x100 is above the range but stays inside the cross-field rule."""
    value = sales["quantity"] * sales["unit_price"] * 100
    return (sales["unit_price"] > 5.01) & (sales["quantity"] > 0) & (value <= 50000)


def inject_cents(sales: pd.DataFrame, idx: np.ndarray) -> None:
    """Multiply the given rows' unit_price by 100 (entered in cents), in place."""
    sales.loc[idx, "unit_price"] = (sales.loc[idx, "unit_price"] * 100).round(2)


def recode_store_channel(ctx: Ctx, sales: pd.DataFrame) -> None:
    """Code the 'store' channel with its variants Store, B&M and 1."""
    mask = sales["channel"] == "store"
    sales.loc[mask, "channel"] = ctx.rng.choice(["Store", "B&M", "1"], int(mask.sum()))


def bad_skus(ctx: Ctx, n: int) -> np.ndarray:
    """SKUs that cannot exist in the master (which uses the SKU00xxx range)."""
    return np.array([f"SKU9{int(v):04d}" for v in ctx.rng.integers(0, 10000, n)])


def add_duplicates(ctx: Ctx, sales: pd.DataFrame, source_idx: np.ndarray) -> pd.DataFrame:
    """Append exact copies of the given rows and shuffle row order."""
    combined = pd.concat([sales, sales.loc[source_idx]], ignore_index=True)
    order = ctx.rng.permutation(len(combined))
    return combined.iloc[order].reset_index(drop=True)


def pick(ctx: Ctx, candidates: np.ndarray, n: int) -> np.ndarray:
    """Choose n distinct candidates, sorted."""
    return np.sort(ctx.rng.choice(candidates, n, replace=False))


# --- renaming and ground truth ----------------------------------------------------------


def synonym(ctx: Ctx, table: str, field: str, name: str) -> str:
    """Return ``name`` after checking it is a declared synonym of the field."""
    spec = ctx.model.tables[table].fields[field]
    assert name in spec.synonyms, f"{name!r} is not a synonym of {spec.qualified_name}"
    return name


def renamed_tables(ctx: Ctx, sales: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """Return copies of sales, products and stores with columns renamed to synonyms."""
    sales_names = {
        "transaction_date": "Txn_Date", "store_id": "Store_No", "sku": "Item_No",
        "quantity": "Units_Sold", "unit_price": "Sell_Price", "channel": "Sales_Channel",
        "promo_flag": "On_Promo",
    }  # fmt: skip
    renames = {
        "sales": sales_names,
        "products": {"sku": "UPC", "category": "Dept", "pack_size": "Case_Pack"},
        "stores": {"store_id": "Store_No", "region": "District"},
    }
    frames = {"sales": sales, "products": ctx.products, "stores": ctx.stores}
    out = {}
    for table, mapping in renames.items():
        checked = {f: synonym(ctx, table, f, s) for f, s in mapping.items()}
        out[table] = frames[table].rename(columns=checked)
    return out


def add_unmatched_columns(ctx: Ctx, sales: pd.DataFrame, columns: list[str]) -> None:
    """Add the requested no-match columns (Cashier_Notes, Legacy_Code) in place."""
    n = len(sales)
    if "Cashier_Notes" in columns:
        sales["Cashier_Notes"] = ctx.rng.choice(CASHIER_NOTES, n)
    if "Legacy_Code" in columns:
        sales["Legacy_Code"] = [f"L{int(v):05d}" for v in ctx.rng.integers(0, 99999, n)]


def field_mapping(
    tables: dict[str, pd.DataFrame], model: DataModel, unmatched: tuple[str, ...] = ()
) -> dict[str, str | None]:
    """Flat source-column mapping across tables; raises if a name maps two ways."""
    mapping: dict[str, str | None] = {}
    for table, frame in tables.items():
        fields = model.tables[table].fields
        for column in frame.columns:
            target = None if column in unmatched else _resolve(column, fields)
            if mapping.get(column, target) != target:
                raise ValueError(f"column {column!r} maps to different fields across tables")
            mapping[column] = target
    return mapping


def _resolve(column: str, fields: dict) -> str:
    for name, spec in fields.items():
        if column == name or column in spec.synonyms:
            return name
    raise ValueError(f"column {column!r} has no field in the data model")


def truth(
    description: str,
    files: dict[str, str],
    rows: dict[str, int],
    mapping: dict[str, str | None],
    missing: list[str] | None = None,
    issues: list[dict[str, Any]] | None = None,
    **extra: Any,
) -> dict[str, Any]:
    """Assemble one fixture's ground_truth.json content."""
    return {
        "description": description,
        "files": files,
        "rows": rows,
        "expected_field_mapping": mapping,
        "expected_missing_required_fields": missing or [],
        "expected_issues": issues or [],
        **extra,
    }


def issue(field: str, issue_type: str, count: int, approx: bool = False) -> dict[str, Any]:
    """One expected issue; approx=True records count_approx (+/-10%)."""
    return {
        "field": field,
        "issue_type": issue_type,
        ("count_approx" if approx else "count"): count,
    }


# --- writers ----------------------------------------------------------------------------


def write_csv(frame: pd.DataFrame, path: Path) -> None:
    """Write a table as CSV. Writes: path."""
    frame.to_csv(path, index=False, lineterminator="\n", float_format="%.2f")


def write_tables(ctx: Ctx, tables: dict[str, pd.DataFrame]) -> tuple[dict, dict]:
    """Write each table as <table>.csv. Returns (files, rows) for the ground truth."""
    files, rows = {}, {}
    for table, frame in tables.items():
        write_csv(frame, ctx.out_dir / f"{table}.csv")
        files[table], rows[table] = f"{table}.csv", len(frame)
    return files, rows


def masters(ctx: Ctx) -> dict[str, pd.DataFrame]:
    """The product and store masters, unchanged."""
    return {"products": ctx.products, "stores": ctx.stores}


def normalize_zip(path: Path) -> None:
    """Rewrite a zip (xlsx) with fixed timestamps so the bytes are deterministic."""
    stamp = FIXED_TIMESTAMP.strftime("%Y-%m-%dT%H:%M:%SZ")
    with zipfile.ZipFile(path) as src:
        entries = [(info.filename, src.read(info)) for info in src.infolist()]
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, data in entries:
            if name == "docProps/core.xml":
                text = re.sub(r"(<dcterms:(created|modified)[^>]*>)[^<]*", rf"\g<1>{stamp}",
                              data.decode())  # fmt: skip
                data = text.encode()
            dst.writestr(zipfile.ZipInfo(name, date_time=ZIP_EPOCH), data, zipfile.ZIP_DEFLATED)
    path.write_bytes(buffer.getvalue())


def write_messy_excel(sales: pd.DataFrame, path: Path) -> None:
    """Two sheets: a summary sheet, then the data sheet with its header on row 3."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Weekly sales export"])
    summary.append(["Generated by POS v2"])
    data = workbook.create_sheet("Sales")
    data.append(["Store sales export"])
    data.append(["Period: 2023-01-01 to 2024-12-31"])
    data.append(list(sales.columns))
    for row in sales.itertuples(index=False):
        values = list(row)
        values[0] = date.fromisoformat(values[0])
        data.append([v.item() if isinstance(v, np.generic) else v for v in values])
    workbook.properties.created = workbook.properties.modified = FIXED_TIMESTAMP
    workbook.save(path)
    normalize_zip(path)


def email_base(subject: str) -> EmailMessage:
    """An email with fixed headers, so the bytes are deterministic."""
    message = EmailMessage()
    message["From"] = "data-team@customer.example"
    message["To"] = "onboarding@forecasting.example"
    message["Subject"] = subject
    message["Date"] = format_datetime(FIXED_TIMESTAMP)
    message["Message-ID"] = f"<{zlib.crc32(subject.encode())}@fixtures.dva>"
    return message


def write_csv_email(frame: pd.DataFrame, path: Path, filename: str) -> None:
    """Write an .eml with the table as a CSV attachment. Writes: path."""
    message = email_base("Sales export attached")
    message.set_content("Hi, please find this week's sales export attached.\n")
    payload = frame.to_csv(index=False, lineterminator="\n", float_format="%.2f").encode()
    message.add_attachment(payload, maintype="text", subtype="csv", filename=filename)
    message.set_boundary(EMAIL_BOUNDARY)
    path.write_bytes(bytes(message))


def write_html_email(frame: pd.DataFrame, path: Path) -> None:
    """Write an .eml with the table as an HTML table in the body. Writes: path."""
    message = email_base("Store master")
    message.set_content("Store master below (HTML table).\n")
    header = "".join(f"<th>{c}</th>" for c in frame.columns)
    body = "".join(
        "<tr>" + "".join(f"<td>{v}</td>" for v in row) + "</tr>"
        for row in frame.itertuples(index=False)
    )
    html = f"<html><body><p>Store master:</p><table><tr>{header}</tr>{body}</table></body></html>"
    message.add_alternative(html, subtype="html")
    message.set_boundary(EMAIL_BOUNDARY)
    path.write_bytes(bytes(message))


# --- fixtures ---------------------------------------------------------------------------


def build_clean(ctx: Ctx) -> dict[str, Any]:
    """Write the clean fixture into ctx.out_dir. Returns its ground truth."""
    tables = {"sales": make_sales(ctx, 2000), **masters(ctx)}
    files, rows = write_tables(ctx, tables)
    return truth("No issues. Every check should pass.", files, rows,
                 field_mapping(tables, ctx.model))  # fmt: skip


def build_renamed_columns(ctx: Ctx) -> dict[str, Any]:
    """Write the renamed_columns fixture into ctx.out_dir. Returns its ground truth."""
    sales = make_sales(ctx, 2000)
    tables = renamed_tables(ctx, sales)
    unmatched = ("Cashier_Notes", "Legacy_Code")
    add_unmatched_columns(ctx, tables["sales"], list(unmatched))
    files, rows = write_tables(ctx, tables)
    issues = [issue(col, "unmapped_column", rows["sales"]) for col in unmatched]
    return truth("Source names from the synonym lists, plus Cashier_Notes and Legacy_Code.",
                 files, rows, field_mapping(tables, ctx.model, unmatched),
                 issues=issues)  # fmt: skip


def build_missing_field(ctx: Ctx) -> dict[str, Any]:
    """Write the missing_field fixture into ctx.out_dir. Returns its ground truth."""
    sales = make_sales(ctx, 2000).drop(columns=["channel"])
    tables = {"sales": sales, **masters(ctx)}
    files, rows = write_tables(ctx, tables)
    return truth("Required field channel is absent from sales.", files, rows,
                 field_mapping(tables, ctx.model), missing=["channel"],
                 issues=[issue("channel", "missing_field", rows["sales"])])  # fmt: skip


def build_unit_mismatch(ctx: Ctx) -> dict[str, Any]:
    """Write the unit_mismatch fixture into ctx.out_dir. Returns its ground truth."""
    case_stores = ["S007", "S019", "S033"]
    sales = make_sales(ctx, 5000, case_store_weights(ctx, case_stores, 0.08))
    affected = inject_cases(sales, case_stores)
    tables = {"sales": sales, **masters(ctx)}
    files, rows = write_tables(ctx, tables)
    return truth(f"Quantities x{CASE_PACK} (cases, not units) in stores {case_stores}.",
                 files, rows, field_mapping(tables, ctx.model),
                 issues=[issue("quantity", "unit_mismatch", affected, approx=True)],
                 case_stores=case_stores)  # fmt: skip


def build_implausible_values(ctx: Ctx) -> dict[str, Any]:
    """Write the implausible_values fixture into ctx.out_dir. Returns its ground truth."""
    sales = make_sales(ctx, 5000)
    idx = pick(ctx, sales.index[price_eligible(sales)].to_numpy(), 37)
    inject_cents(sales, idx)
    tables = {"sales": sales, **masters(ctx)}
    files, rows = write_tables(ctx, tables)
    return truth("37 unit prices multiplied by 100 (entered in cents).", files, rows,
                 field_mapping(tables, ctx.model),
                 issues=[issue("unit_price", "implausible_value", 37)])  # fmt: skip


def build_category_variants(ctx: Ctx) -> dict[str, Any]:
    """Write the category_variants fixture into ctx.out_dir. Returns its ground truth."""
    sales = make_sales(ctx, 3000)
    recode_store_channel(ctx, sales)
    sales.loc[pick(ctx, sales.index.to_numpy(), 25), "channel"] = "ZZ"
    tables = {"sales": sales, **masters(ctx)}
    files, rows = write_tables(ctx, tables)
    return truth("channel coded as Store, B&M and 1, plus 25 rows with unknown code ZZ.",
                 files, rows, field_mapping(tables, ctx.model),
                 issues=[issue("channel", "unmapped_category", 25)])  # fmt: skip


def build_join_failure(ctx: Ctx) -> dict[str, Any]:
    """Write the join_failure fixture into ctx.out_dir. Returns its ground truth."""
    sales = make_sales(ctx, 4000)
    sales.loc[pick(ctx, sales.index.to_numpy(), 412), "sku"] = bad_skus(ctx, 412)
    tables = {"sales": sales, **masters(ctx)}
    files, rows = write_tables(ctx, tables)
    return truth("412 sales rows have a SKU missing from the 300-product master.", files,
                 rows, field_mapping(tables, ctx.model),
                 issues=[issue("sku", "join_failure", 412)])  # fmt: skip


def build_type_errors(ctx: Ctx) -> dict[str, Any]:
    """Write the type_errors fixture into ctx.out_dir. Returns its ground truth."""
    sales = make_sales(ctx, 2000)
    formats = ["%Y-%m-%d", "%m/%d/%Y", "%d-%b-%Y", "%Y%m%d"]
    fmt_idx = ctx.rng.choice(len(formats), len(sales), p=[0.4, 0.2, 0.2, 0.2])
    sales["transaction_date"] = [
        date.fromisoformat(d).strftime(formats[i])
        for d, i in zip(sales["transaction_date"], fmt_idx, strict=True)
    ]
    style = ctx.rng.choice(3, len(sales), p=[0.4, 0.4, 0.2])  # plain, "$4.20", "4,20"
    styles = ["{:.2f}", "${:,.2f}", "{:.2f}"]
    sales["unit_price"] = [
        styles[s].format(p).replace(".", ",") if s == 2 else styles[s].format(p)
        for p, s in zip(sales["unit_price"], style, strict=True)
    ]
    tables = {"sales": sales, **masters(ctx)}
    files, rows = write_tables(ctx, tables)
    issues = [issue("transaction_date", "type_mismatch", int((fmt_idx != 0).sum())),
              issue("unit_price", "type_mismatch", int((style != 0).sum()))]  # fmt: skip
    return truth("Dates as text in mixed formats; prices with $ or a decimal comma.", files,
                 rows, field_mapping(tables, ctx.model), issues=issues)  # fmt: skip


def build_duplicates(ctx: Ctx) -> dict[str, Any]:
    """Write the duplicates fixture into ctx.out_dir. Returns its ground truth."""
    sales = make_sales(ctx, 1950)
    sales = add_duplicates(ctx, sales, pick(ctx, sales.index.to_numpy(), 50))
    tables = {"sales": sales, **masters(ctx)}
    files, rows = write_tables(ctx, tables)
    return truth("50 exact duplicate rows.", files, rows, field_mapping(tables, ctx.model),
                 issues=[issue("*", "duplicate_rows", 50)])  # fmt: skip


def build_excel_messy(ctx: Ctx) -> dict[str, Any]:
    """Write the excel_messy fixture into ctx.out_dir. Returns its ground truth."""
    sales = make_sales(ctx, 3000)
    write_messy_excel(sales, ctx.out_dir / "sales.xlsx")
    files, rows = write_tables(ctx, masters(ctx))
    tables = {"sales": sales, **masters(ctx)}
    return truth("Sales in .xlsx: two sheets, data on sheet 'Sales' with the header on row 3.",
                 {"sales": "sales.xlsx", **files}, {"sales": len(sales), **rows},
                 field_mapping(tables, ctx.model),
                 excel={"file": "sales.xlsx", "sheet": "Sales", "header_row": 3})  # fmt: skip


def build_email_table(ctx: Ctx) -> dict[str, Any]:
    """Write the email_table fixture into ctx.out_dir. Returns its ground truth."""
    sales = make_sales(ctx, 500)
    write_csv_email(sales, ctx.out_dir / "sales.eml", "sales_export.csv")
    write_html_email(ctx.stores, ctx.out_dir / "stores.eml")
    write_csv(ctx.products, ctx.out_dir / "products.csv")
    tables = {"sales": sales, **masters(ctx)}
    return truth("sales.eml has a CSV attachment; stores.eml has an HTML table in the body.",
                 {"sales": "sales.eml", "products": "products.csv", "stores": "stores.eml"},
                 {name: len(frame) for name, frame in tables.items()},
                 field_mapping(tables, ctx.model))  # fmt: skip


def build_kitchen_sink(ctx: Ctx) -> dict[str, Any]:
    """Write the kitchen_sink fixture into ctx.out_dir. Returns its ground truth."""
    case_stores = ["S004", "S021", "S038"]
    sales = make_sales(ctx, 9950, case_store_weights(ctx, case_stores, 0.08))
    pool = ctx.rng.permutation(sales.index[~sales["store_id"].isin(case_stores)].to_numpy())
    eligible = price_eligible(sales).to_numpy()
    cents_idx = np.sort([i for i in pool if eligible[i]][:37])
    rest = np.array([i for i in pool if i not in set(cents_idx)])
    zz_idx, sku_idx, dup_idx = np.sort(rest[:25]), np.sort(rest[25:175]), rest[175:225]
    affected = inject_cases(sales, case_stores)
    inject_cents(sales, cents_idx)
    recode_store_channel(ctx, sales)
    sales.loc[zz_idx, "channel"] = "ZZ"
    sales.loc[sku_idx, "sku"] = bad_skus(ctx, len(sku_idx))
    add_unmatched_columns(ctx, sales, ["Cashier_Notes"])  # before duplicating: exact copies
    sales = add_duplicates(ctx, sales, dup_idx)
    tables = renamed_tables(ctx, sales)
    files, rows = write_tables(ctx, tables)
    return truth("Renamed columns + Cashier_Notes, cases in 3 stores, 37 cent prices, "
                 "25 ZZ channel codes, 150 unknown SKUs, 50 duplicates.",
                 files, rows, field_mapping(tables, ctx.model, ("Cashier_Notes",)),
                 issues=kitchen_sink_issues(rows["sales"], affected),
                 case_stores=case_stores)  # fmt: skip


def kitchen_sink_issues(sales_rows: int, case_rows: int) -> list[dict[str, Any]]:
    """Expected issues for kitchen_sink."""
    return [
        issue("Cashier_Notes", "unmapped_column", sales_rows),
        issue("quantity", "unit_mismatch", case_rows, approx=True),
        issue("unit_price", "implausible_value", 37),
        issue("channel", "unmapped_category", 25),
        issue("sku", "join_failure", 150),
        issue("*", "duplicate_rows", 50),
    ]


BUILDERS: dict[str, Callable[[Ctx], dict[str, Any]]] = {
    "clean": build_clean,
    "renamed_columns": build_renamed_columns,
    "missing_field": build_missing_field,
    "unit_mismatch": build_unit_mismatch,
    "implausible_values": build_implausible_values,
    "category_variants": build_category_variants,
    "join_failure": build_join_failure,
    "type_errors": build_type_errors,
    "duplicates": build_duplicates,
    "excel_messy": build_excel_messy,
    "email_table": build_email_table,
    "kitchen_sink": build_kitchen_sink,
}


# --- entry points -----------------------------------------------------------------------


def make_context(name: str, out_dir: Path, model: DataModel) -> Ctx:
    """A per-fixture RNG keyed by name, so fixtures are independent of each other."""
    rng = np.random.default_rng([SEED, zlib.crc32(name.encode())])
    stores = make_stores(rng)
    products, prices = make_products(rng)
    return Ctx(name, rng, out_dir, model, stores, products, prices)


def generate_all(out_root: Path, model: DataModel) -> dict[str, dict[str, Any]]:
    """Generate every fixture under out_root.

    Writes: out_root/<fixture>/ (replaced if present) with data files and ground_truth.json.
    Returns: the ground truth per fixture.
    """
    results = {}
    for name, build in BUILDERS.items():
        out_dir = out_root / name
        shutil.rmtree(out_dir, ignore_errors=True)
        out_dir.mkdir(parents=True)
        ground_truth = build(make_context(name, out_dir, model))
        text = json.dumps(ground_truth, indent=2, sort_keys=False) + "\n"
        (out_dir / "ground_truth.json").write_text(text, encoding="utf-8")
        results[name] = ground_truth
    return results


def main() -> None:
    """CLI entry point. Writes: the fixtures folder from settings, or --out."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=None, help="output folder")
    args = parser.parse_args()
    settings = load_settings()
    model = make_provider(settings).get_model(settings.data_model.default)
    out_root = args.out or settings.paths.fixtures_dir
    for name, gt in generate_all(out_root, model).items():
        print(f"{name:20s} rows={gt['rows']} issues={len(gt['expected_issues'])}")


if __name__ == "__main__":
    main()
