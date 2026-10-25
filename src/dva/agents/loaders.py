"""Load customer files into raw tables: CSV, Excel, and email (.eml).

Loaders never raise for bad customer data. A file that cannot be read becomes a
LoadIssue, which parsing turns into a finding. Cell values stay as loaded (strings for
text sources, native types for Excel) so parsing can report type problems.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from collections.abc import Iterable
from email import message_from_bytes, policy
from email.message import EmailMessage
from html.parser import HTMLParser
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, ConfigDict

CSV_SUFFIXES = {".csv", ".tsv", ".txt"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
EMAIL_SUFFIXES = {".eml"}
HEADER_SCAN_ROWS = 20
MIN_HEADER_MATCHES = 2
ENCODINGS = ("utf-8-sig", "cp1252")
READ_ERRORS = (OSError, ValueError, UnicodeError, csv.Error, zipfile.BadZipFile,
               InvalidFileException, KeyError)  # fmt: skip


class RawTable(BaseModel):
    """One table found in a customer file, before mapping."""

    model_config = ConfigDict(arbitrary_types_allowed=True)

    source: str  # file name, plus "#sheet" or "#attachment" where relevant
    frame: pd.DataFrame
    header_row: int = 1  # 1-based row of the header in the source
    sheet: str | None = None


class LoadIssue(BaseModel):
    """A file, sheet, or part that could not be loaded."""

    source: str
    detail: str


class LoadResult(BaseModel):
    tables: list[RawTable] = []
    issues: list[LoadIssue] = []


def normalize_name(name: object) -> str:
    """Lowercase and drop everything but letters and digits: 'Unit Price' -> 'unitprice'."""
    return re.sub(r"[^a-z0-9]", "", str(name).lower())


def load_files(paths: Iterable[Path], vocabulary: set[str]) -> LoadResult:
    """Load every file by extension.

    ``vocabulary`` holds normalized field names and synonyms; it is used to find header
    rows and to pick data sheets. Reads: the files. Writes nothing.
    """
    result = LoadResult()
    for path in paths:
        path = Path(path)
        suffix = path.suffix.lower()
        try:
            if suffix in CSV_SUFFIXES:
                part = LoadResult(tables=[load_csv_bytes(path.read_bytes(), path.name)])
            elif suffix in EXCEL_SUFFIXES:
                part = load_excel(path, vocabulary)
            elif suffix in EMAIL_SUFFIXES:
                part = load_email(path, vocabulary)
            else:
                issue = LoadIssue(source=path.name, detail=f"unsupported file type {suffix!r}")
                part = LoadResult(issues=[issue])
        except READ_ERRORS as exc:
            part = LoadResult(issues=[LoadIssue(source=path.name, detail=f"unreadable: {exc}")])
        result.tables += part.tables
        result.issues += part.issues
    return result


# --- CSV --------------------------------------------------------------------------------


def load_csv_bytes(data: bytes, source: str) -> RawTable:
    """Parse delimited text with every cell as a string; blanks become ''."""
    text = _decode(data)
    delimiter = _sniff_delimiter(text)
    frame = pd.read_csv(io.StringIO(text), sep=delimiter, dtype=str, keep_default_na=False)
    frame.columns = [str(c).strip() for c in frame.columns]
    return RawTable(source=source, frame=frame)


def _decode(data: bytes) -> str:
    for encoding in ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnicodeError("not valid UTF-8 or Windows-1252 text")


def _sniff_delimiter(text: str) -> str:
    sample = "\n".join(text.splitlines()[:20])
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


# --- Excel ------------------------------------------------------------------------------


def load_excel(path: Path, vocabulary: set[str]) -> LoadResult:
    """Load every sheet that has a recognizable header row (sheet selection).

    A header row is the row in the first 20 with the most cells matching the
    vocabulary; a sheet with fewer than 2 matches is skipped.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    result = LoadResult()
    try:
        for sheet in workbook.worksheets:
            rows = [list(r) for r in sheet.iter_rows(values_only=True)]
            table = rows_to_table(rows, vocabulary, f"{path.name}#{sheet.title}")
            if table is not None:
                result.tables.append(table.model_copy(update={"sheet": sheet.title}))
    finally:
        workbook.close()
    if not result.tables:
        result.issues.append(LoadIssue(source=path.name, detail="no sheet has a header row "
                                       "matching the data model"))  # fmt: skip
    return result


def rows_to_table(rows: list[list], vocabulary: set[str], source: str) -> RawTable | None:
    """Build a table from raw rows using the detected header row, or None if none is found."""
    header_idx = detect_header_row(rows, vocabulary)
    if header_idx is None:
        return None
    header = [str(c).strip() if c is not None else "" for c in rows[header_idx]]
    keep = [i for i, name in enumerate(header) if name]
    body = [
        [row[i] if i < len(row) else None for i in keep]
        for row in rows[header_idx + 1 :]
        if any(v not in (None, "") for v in row)
    ]
    frame = pd.DataFrame(body, columns=[header[i] for i in keep], dtype=object)
    return RawTable(source=source, frame=frame, header_row=header_idx + 1)


def detect_header_row(rows: list[list], vocabulary: set[str]) -> int | None:
    """Return the 0-based index of the best header row in the first rows, or None."""
    best_idx, best_score = None, MIN_HEADER_MATCHES - 1
    for idx, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        score = sum(1 for cell in row if cell is not None and normalize_name(cell) in vocabulary)
        if score > best_score:
            best_idx, best_score = idx, score
    return best_idx


# --- Email ------------------------------------------------------------------------------


def load_email(path: Path, vocabulary: set[str]) -> LoadResult:
    """Load CSV attachments and HTML tables in the body of an .eml file."""
    message = message_from_bytes(path.read_bytes(), policy=policy.default)
    assert isinstance(message, EmailMessage)
    result = LoadResult()
    for part in message.iter_attachments():
        filename = part.get_filename() or "attachment"
        if Path(filename).suffix.lower() in CSV_SUFFIXES or part.get_content_type() == "text/csv":
            payload = part.get_payload(decode=True) or b""
            result.tables.append(load_csv_bytes(payload, f"{path.name}#{filename}"))
    body = message.get_body(preferencelist=("html",))
    if body is not None:
        for n, rows in enumerate(parse_html_tables(body.get_content()), start=1):
            table = rows_to_table(rows, vocabulary, f"{path.name}#html-table-{n}")
            if table is not None:
                result.tables.append(table)
    if not result.tables:
        result.issues.append(LoadIssue(source=path.name, detail="no CSV attachment or HTML "
                                       "table matching the data model"))  # fmt: skip
    return result


class _TableParser(HTMLParser):
    """Collect the text of <th>/<td> cells per row for every <table>."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list) -> None:
        if tag == "table":
            self.tables.append([])
        elif tag == "tr" and self.tables:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.tables[-1].append(self._row)
            self._row = None

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)


def parse_html_tables(html: str) -> list[list[list[str]]]:
    """Return each HTML table as a list of rows of cell text (stdlib parser only)."""
    parser = _TableParser()
    parser.feed(html)
    parser.close()
    return [table for table in parser.tables if table]
